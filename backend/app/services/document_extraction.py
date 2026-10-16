from __future__ import annotations

from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from typing import Any

import extract_msg
from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.pdf_vision import extract_scanned_pdf_with_openai
from app.services.image_vision import extract_image_with_openai


@dataclass
class ExtractedSection:
    text: str
    page_number: int | None = None
    sheet_name: str | None = None
    email_section: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class ExtractedAttachment:
    """Raw email attachment that can be recursively extracted/indexed.

    Attachments remain part of the immutable parent email document in storage.
    The indexing pipeline extracts supported attachment contents into dedicated
    document chunks and preserves the parent/attachment provenance in metadata.
    """

    file_name: str
    content: bytes
    content_type: str | None = None


@dataclass
class ExtractedDocument:
    source_type: str
    sections: list[ExtractedSection]
    metadata: dict[str, Any] = field(default_factory=dict)
    attachments: list[ExtractedAttachment] = field(default_factory=list)


def extract_pdf(
    content: bytes,
    *,
    file_name: str = "document.pdf",
) -> ExtractedDocument:
    reader = PdfReader(BytesIO(content))
    sections: list[ExtractedSection] = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(
                ExtractedSection(
                    text=text,
                    page_number=page_number,
                    metadata={"vision_extraction": False},
                )
            )

    extracted_character_count = sum(len(section.text) for section in sections)
    if extracted_character_count >= 50:
        return ExtractedDocument(
            "pdf",
            sections,
            {
                "page_count": len(reader.pages),
                "vision_extraction": False,
            },
        )

    vision_result = extract_scanned_pdf_with_openai(
        content,
        file_name=file_name,
    )
    vision_sections = [
        ExtractedSection(
            text=page.text.strip(),
            page_number=page.page_number,
            metadata={
                "vision_extraction": True,
                "extraction_provider": "openai",
            },
        )
        for page in vision_result.pages
        if page.text.strip()
    ]

    return ExtractedDocument(
        "pdf_vision",
        vision_sections,
        {
            "page_count": len(reader.pages),
            "vision_extraction": True,
            "extraction_provider": "openai",
        },
    )


def extract_docx(content: bytes) -> ExtractedDocument:
    document = WordDocument(BytesIO(content))
    sections: list[ExtractedSection] = []
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if paragraphs:
        sections.append(ExtractedSection("\n\n".join(paragraphs)))
    for index, table in enumerate(document.tables):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        table_text = "\n".join(row for row in rows if row.strip()).strip()
        if table_text:
            sections.append(
                ExtractedSection(table_text, metadata={"table_index": index})
            )
    if not sections:
        raise ValueError("Das DOCX-Dokument enthält keinen extrahierbaren Text.")
    return ExtractedDocument("docx", sections)


def extract_xlsx(content: bytes) -> ExtractedDocument:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections: list[ExtractedSection] = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append(" | ".join(values))
        if rows:
            sections.append(
                ExtractedSection("\n".join(rows), sheet_name=worksheet.title)
            )
    if not sections:
        raise ValueError("Die XLSX-Datei enthält keine verwertbaren Zellwerte.")
    return ExtractedDocument("xlsx", sections, {"sheet_count": len(workbook.sheetnames)})


def _decode_text_attachment(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace").strip()


def extract_text(content: bytes, *, source_type: str = "text") -> ExtractedDocument:
    text = _decode_text_attachment(content)
    if not text:
        raise ValueError("Die Textdatei enthält keinen extrahierbaren Inhalt.")
    return ExtractedDocument(source_type, [ExtractedSection(text)])


def extract_eml(content: bytes) -> ExtractedDocument:
    message = BytesParser(policy=policy.default).parsebytes(content)
    body_parts: list[str] = []
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_type() == "text/plain" and not part.get_filename():
                try:
                    value = part.get_content()
                    if value and value.strip():
                        body_parts.append(value.strip())
                except Exception:
                    continue
    else:
        try:
            value = message.get_content()
            if value and value.strip():
                body_parts.append(value.strip())
        except Exception:
            pass

    sender = str(message.get("From", "") or "").strip() or None
    recipient = str(message.get("To", "") or "").strip() or None
    subject = str(message.get("Subject", "") or "").strip() or None
    document_date = str(message.get("Date", "") or "").strip() or None

    header = "\n".join(
        [
            f"Von: {sender or ''}",
            f"An: {recipient or ''}",
            f"Betreff: {subject or ''}",
            f"Datum: {document_date or ''}",
        ]
    )

    attachments: list[ExtractedAttachment] = []
    attachment_names: list[str] = []
    for index, part in enumerate(message.iter_attachments(), start=1):
        file_name = Path(part.get_filename() or f"attachment-{index}").name
        try:
            payload = part.get_payload(decode=True)
        except Exception:
            payload = None
        if not isinstance(payload, (bytes, bytearray)):
            continue
        attachment_names.append(file_name)
        attachments.append(
            ExtractedAttachment(
                file_name=file_name,
                content=bytes(payload),
                content_type=part.get_content_type(),
            )
        )

    text = f"{header}\n\n" + "\n\n".join(body_parts)
    if attachment_names:
        text += "\n\nAnhänge: " + ", ".join(attachment_names)

    return ExtractedDocument(
        "eml",
        [ExtractedSection(text.strip(), email_section="message")],
        {
            "attachments": attachment_names,
            "attachment_count": len(attachment_names),
            "sender": sender,
            "recipient": recipient,
            "subject": subject,
            "document_date": document_date,
        },
        attachments,
    )


def _msg_attachment_bytes(attachment: Any) -> bytes | None:
    """Best-effort byte extraction for extract-msg attachment variants."""
    try:
        data = getattr(attachment, "data", None)
        if callable(data):
            data = data()
        if isinstance(data, bytearray):
            return bytes(data)
        if isinstance(data, bytes):
            return data
    except Exception:
        return None
    return None


def extract_msg_document(content: bytes) -> ExtractedDocument:
    message = extract_msg.Message(BytesIO(content))
    try:
        sender = str(message.sender or "").strip() or None
        recipient = str(message.to or "").strip() or None
        subject = str(message.subject or "").strip() or None
        document_date = str(message.date or "").strip() or None
        text = "\n".join(
            [
                f"Von: {sender or ''}",
                f"An: {recipient or ''}",
                f"Betreff: {subject or ''}",
                f"Datum: {document_date or ''}",
                "",
                message.body or "",
            ]
        ).strip()
        if not text:
            raise ValueError("Die MSG-Datei enthält keinen extrahierbaren Inhalt.")

        attachments: list[ExtractedAttachment] = []
        attachment_names: list[str] = []
        for index, attachment in enumerate(getattr(message, "attachments", []) or [], start=1):
            file_name = (
                getattr(attachment, "longFilename", None)
                or getattr(attachment, "shortFilename", None)
                or f"attachment-{index}"
            )
            file_name = Path(str(file_name)).name
            payload = _msg_attachment_bytes(attachment)
            if payload is None:
                continue
            attachment_names.append(file_name)
            attachments.append(
                ExtractedAttachment(
                    file_name=file_name,
                    content=payload,
                    content_type=getattr(attachment, "mimetype", None),
                )
            )

        if attachment_names:
            text += "\n\nAnhänge: " + ", ".join(attachment_names)

        return ExtractedDocument(
            "msg",
            [ExtractedSection(text, email_section="message")],
            {
                "attachments": attachment_names,
                "attachment_count": len(attachment_names),
                "sender": sender,
                "recipient": recipient,
                "subject": subject,
                "document_date": document_date,
            },
            attachments,
        )
    finally:
        try:
            message.close()
        except Exception:
            pass



def extract_image(content: bytes, *, file_name: str) -> ExtractedDocument:
    """OCR a case-relevant standalone/attached image via OpenAI vision."""
    text = extract_image_with_openai(content, file_name=file_name)
    metadata = {
        "ocr": True,
        "vision_extraction": True,
        "extraction_provider": "openai",
    }
    return ExtractedDocument(
        "image_ocr",
        [ExtractedSection(text=text, metadata=dict(metadata))],
        metadata,
    )


def extract_document(*, file_name: str, content: bytes) -> ExtractedDocument:
    extension = Path(file_name).suffix.lower()
    if extension == ".pdf":
        return extract_pdf(content, file_name=file_name)
    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        return extract_image(content, file_name=file_name)
    if extension == ".docx":
        return extract_docx(content)
    if extension == ".xlsx":
        return extract_xlsx(content)
    if extension == ".eml":
        return extract_eml(content)
    if extension == ".msg":
        return extract_msg_document(content)
    if extension == ".txt":
        return extract_text(content, source_type="txt")
    if extension == ".csv":
        return extract_text(content, source_type="csv")
    if extension in {".doc", ".xls"}:
        raise ValueError(
            "Legacy-DOC/XLS muss vor der KI-Verarbeitung in DOCX/XLSX umgewandelt werden."
        )
    raise ValueError(f"Das Format {extension or 'unbekannt'} wird nicht unterstützt.")
